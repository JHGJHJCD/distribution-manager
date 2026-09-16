# -*- coding: utf-8 -*-
"""v3.52 נתמך חגים — assert + screenshot: recipient form (mark + per-holiday row),
recipients-tab filter 'נתמך חגים' (+ holiday sub-combo), manual-add picker filter,
and the custom-filter dialog's holiday gate through db.get_filtered_list."""
import sys, os, tempfile
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from PyQt6.QtWidgets import QApplication
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QFont

import database as db
d = tempfile.mkdtemp(); db.DB_PATH = os.path.join(d, "x.db"); db.BACKUP_DIR = os.path.join(d, "b")
db.init_db()
seed = [
    ("כהן ישראל",  4, "שבועי",  1, ""),
    ("לוי משה",    3, "חד-פעמי", 1, "פסח,סוכות"),
    ("מזרחי דוד",  2, "חד-פעמי", 0, ""),
    ("פרץ יוסף",   4, "חודשי",  0, ""),
]
ids = {}
for name, pr, freq, hs, hl in seed:
    ids[name] = db.add_recipient({"full_name": name, "status": "פעיל", "priority": pr,
                                  "frequency": freq, "phone1": "0500000000",
                                  "income": "1500", "children_total": 3,
                                  "holiday_support": hs, "holidays": hl})

# DB-level: migration + round-trip + filtered list gate
r = db.get_recipient(ids["לוי משה"])
assert r["holiday_support"] == 1 and r["holidays"] == "פסח,סוכות", r
db.set_setting("available_products", "0")
names = lambda rows: sorted(x["full_name"] for x in rows)
assert names(db.get_filtered_list({"holiday": "*"})) == ["כהן ישראל", "לוי משה"]
assert names(db.get_filtered_list({"holiday": "חנוכה"})) == ["כהן ישראל"]
assert names(db.get_filtered_list({"holiday": "פסח"})) == ["כהן ישראל", "לוי משה"]
assert len(db.get_filtered_list({})) == 4

app = QApplication(sys.argv)
app.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
try:
    from qt_material import apply_stylesheet
    from styles import EXTRA_QSS, QT_MATERIAL_EXTRA
    apply_stylesheet(app, theme="light_teal.xml", invert_secondary=True, extra=QT_MATERIAL_EXTRA)
    app.setStyleSheet(app.styleSheet() + EXTRA_QSS)
except Exception:
    pass
app.setFont(QFont("Segoe UI", 11))
OUT = os.path.join(os.path.dirname(__file__), "_shots")
os.makedirs(OUT, exist_ok=True)

# 1. recipient form
from tabs.recipients import RecipientDialog, RecipientsTab
dlg = RecipientDialog(None, db.get_recipient(ids["לוי משה"]))
dlg.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen, True)
dlg.resize(700, 900); dlg.show(); app.processEvents()
assert dlg.f_holiday.isChecked()
assert dlg._form1.isRowVisible(dlg._holiday_row)
assert dlg.f_holiday_boxes["פסח"].isChecked() and not dlg.f_holiday_boxes["חנוכה"].isChecked()
data = dlg.get_data()
assert data["holiday_support"] == 1 and data["holidays"] == "סוכות,פסח", data
dlg.f_holiday.setChecked(False); app.processEvents()
assert not dlg._form1.isRowVisible(dlg._holiday_row)
assert dlg.get_data()["holidays"] == "" and dlg.get_data()["holiday_support"] == 0
dlg.f_holiday.setChecked(True); app.processEvents()
dlg.grab().save(os.path.join(OUT, "holidays_form.png"))

# 2. recipients tab filter
tab = RecipientsTab()
tab.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen, True)
tab.resize(1400, 600); tab.show(); app.processEvents(); tab.refresh()
assert not tab.holiday_filter.isEnabled()
assert tab.table.rowCount() == 4
tab.priority_filter.setCurrentText("נתמך חגים"); app.processEvents()
assert tab.holiday_filter.isEnabled()
assert tab.table.rowCount() == 2, tab.table.rowCount()
tab.holiday_filter.setCurrentText("חנוכה"); app.processEvents()
assert tab.table.rowCount() == 1 and tab.table.item(0, 1).text() == "ישראל"
tab.holiday_filter.setCurrentText("פסח"); app.processEvents()
assert tab.table.rowCount() == 2
col = tab.table.horizontalHeaderItem(4).text()
assert col == "נתמך חגים", col
assert tab.table.horizontalHeaderItem(14).text() == "סטטוס"
tab.grab().save(os.path.join(OUT, "holidays_tab.png"))
tab.priority_filter.setCurrentText("קבוע"); app.processEvents()
assert tab.table.rowCount() == 2 and not tab.holiday_filter.isEnabled()

# 3. manual-add picker
from tabs.group_update import _ManualAddDialog, FilterCriteriaDialog
m = _ManualAddDialog()
m.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen, True); m.show(); app.processEvents()
m._prio_filter.setCurrentText("נתמך חגים"); app.processEvents()
assert m._holiday_filter.isEnabled() and m._table.rowCount() == 2
m._holiday_filter.setCurrentText("שבועות"); app.processEvents()
assert m._table.rowCount() == 1
m._prio_filter.setCurrentText("כולם"); app.processEvents()
assert not m._holiday_filter.isEnabled() and m._table.rowCount() == 4

# 4. criteria dialog round-trip
f = FilterCriteriaDialog(None, {"holiday": "פסח"})
f.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen, True); f.show(); app.processEvents()
assert f.holiday_combo.currentText() == "נתמכי פסח"
assert f.get_criteria()["holiday"] == "פסח"
f.grab().save(os.path.join(OUT, "holidays_criteria.png"))
f._clear(); assert f.get_criteria()["holiday"] == ""

# 5. Excel export/import round-trip of the column
from utils import excel_utils as xl
import openpyxl
db.set_setting("export_dir_recipients", d)
path = xl.export_recipients_to_excel(db.get_all_recipients())
wb = openpyxl.load_workbook(path); ws = wb.active
hrow = next(r for r in range(1, ws.max_row + 1)
            if "שם מלא" in [c.value for c in ws[r]])
hdr = [c.value for c in ws[hrow]]
assert "נתמך חגים" in hdr, hdr
ci = hdr.index("נתמך חגים")
vals = {ws.cell(r, hdr.index("שם מלא") + 1).value: ws.cell(r, ci + 1).value
        for r in range(hrow + 1, ws.max_row + 1)}
assert vals["לוי משה"] == "סוכות, פסח" and vals["כהן ישראל"] == "כל החגים" and not vals["מזרחי דוד"], vals
back = xl.import_from_excel(path)
bl = {r["full_name"]: r for r in back}
assert bl["לוי משה"]["holiday_support"] == 1 and bl["לוי משה"]["holidays"] == "סוכות,פסח", bl["לוי משה"]
assert bl["כהן ישראל"]["holidays"] == "" and bl["כהן ישראל"]["holiday_support"] == 1
assert bl["מזרחי דוד"]["holiday_support"] == 0
print("OK holidays: form / tab filter / picker / criteria / excel")
