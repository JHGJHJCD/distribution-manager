# -*- coding: utf-8 -*-
"""משימה 6 (6/10/2026): סטטוס המקבל — רק "פעיל" / "מושהה". "הסתיים" ירד.
כל ערך "הסתיים" (כרטיס, אקסל, סנכרון, מקבל שנמחק ושוחזר, DB ישן) הופך ל"מושהה";
ההמרה בהפעלה חד-פעמית, אחרי גיבוי-ביטחון, בלי לגעת ב-updated_at ובלי שינוי שם עמודה.
קמפיין צינתוק עם סטטוס 'done' ("הסתיים" בתצוגה) — לא מושפע.
offscreen, DB זמני."""
import os, re, sys, glob, sqlite3, tempfile
os.environ["QT_QPA_PLATFORM"] = "offscreen"; os.environ["PYTHONUTF8"] = "1"
sys.stdout.reconfigure(encoding="utf-8"); sys.path.insert(0, ".")
import database as db
db.DB_PATH = tempfile.mkstemp(suffix=".db")[1]; db.BACKUP_DIR = tempfile.mkdtemp(); db.init_db()
from PyQt6.QtWidgets import QApplication, QMessageBox
for _m in ("information", "warning", "critical"):
    setattr(QMessageBox, _m, staticmethod(lambda *a, **k: None))
app = QApplication(sys.argv)

fails = []
def ok(name, cond, extra=""):
    print(("  OK  " if cond else "  ✗   ") + name + (f"  [{extra}]" if extra else ""))
    if not cond:
        fails.append(name)

def raw_status(rid):
    with db.get_connection() as c:
        return c.execute("SELECT status FROM recipients WHERE id=?", (rid,)).fetchone()["status"]

ENDED = "הסתיים"
SUSP = "מושהה"

# ── S1: כתיבה דרך הקוד — מקבל חדש / עריכה ───────────────────────────────────────
a = db.add_recipient({"full_name": "כהן אחד", "status": ENDED, "frequency": "שבועי", "priority": 4})
ok("S1a add_recipient 'הסתיים' → 'מושהה'", raw_status(a) == SUSP, raw_status(a))
b = db.add_recipient({"full_name": "לוי שניים", "frequency": "שבועי", "priority": 4})
ok("S1b no status → 'פעיל' (unchanged)", raw_status(b) == "פעיל", raw_status(b))
db.update_recipient(b, {"status": ENDED})
ok("S1c update_recipient 'הסתיים' → 'מושהה'", raw_status(b) == SUSP, raw_status(b))
db.update_recipient(b, {"status": "פעיל"})
ok("S1d back to 'פעיל' works", raw_status(b) == "פעיל")
ok("S1e normalize_status maps only the legacy value",
   db.normalize_status(ENDED) == SUSP and db.normalize_status("פעיל") == "פעיל"
   and db.normalize_status(SUSP) == SUSP and db.normalize_status("") == "פעיל"
   and db.normalize_status(None) == "פעיל" and db.normalize_status(" הסתיים ") == SUSP)

# ── S2: ייבוא אקסל (ייצוא של התוכנה + החלפה/מיזוג) ──────────────────────────────
import openpyxl
from utils import excel_utils
_x = tempfile.mkstemp(suffix=".xlsx")[1]
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["שם מלא", "סטטוס", "תדירות"])
ws.append(["אקסל ישן", ENDED, "שבועי"])
ws.append(["אקסל חי", "פעיל", "שבועי"])
wb.save(_x)
rows = excel_utils.import_from_excel(_x)
st = {r["full_name"]: r["status"] for r in rows}
ok("S2a Excel parse: 'הסתיים' cell → 'מושהה'", st.get("אקסל ישן") == SUSP, str(st))
ok("S2b Excel parse: 'פעיל' stays", st.get("אקסל חי") == "פעיל", str(st))
raw_rows = [{"full_name": "ישיר ישן", "status": ENDED}, {"full_name": "ישיר חדש", "status": "פעיל"}]
n = db.bulk_insert_recipients([dict(r) for r in raw_rows])
with db.get_connection() as c:
    got = {r["full_name"]: r["status"] for r in c.execute(
        "SELECT full_name, status FROM recipients WHERE full_name LIKE 'ישיר%'")}
ok("S2c bulk_insert (replace import) normalizes", got.get("ישיר ישן") == SUSP, str(got))
db.import_recipients_from_list([{"full_name": "מיזוג ישן", "status": ENDED}])
with db.get_connection() as c:
    r = c.execute("SELECT status FROM recipients WHERE full_name='מיזוג ישן'").fetchone()
ok("S2d merge import of a new row normalizes", r and r["status"] == SUSP, str(r and r["status"]))
# הייצוא כותב רק ערכים חוקיים
db.set_setting("export_dir_recipients", tempfile.mkdtemp())   # not the real Downloads
try:
    _out = excel_utils.export_recipients_to_excel(db.get_all_recipients())
    wb2 = openpyxl.load_workbook(_out); ws2 = wb2.active
    hr = next(r for r in range(1, 6) if "סטטוס" in [c.value for c in ws2[r]])
    ci = [c.value for c in ws2[hr]].index("סטטוס")
    vals = {ws2.cell(r, ci + 1).value for r in range(hr + 1, ws2.max_row + 1)} - {None}
    ok("S2e Excel export: statuses ⊆ {פעיל, מושהה}", vals <= {"פעיל", SUSP}, str(vals))
except Exception as e:                                        # noqa: BLE001
    ok("S2e Excel export runs", False, repr(e))

# ── S3: סנכרון — כרטיס מהמחשב השני עם 'הסתיים' ──────────────────────────────────
from utils import sync
card = {k: "" for k in db._RECIPIENT_FIELDS}
card.update({"full_name": "מרחוק חדש", "status": ENDED, "frequency": "שבועי",
             "updated_at": "2030-01-01T00:00:00Z"})
with db.get_connection() as conn:
    sync._apply_rec_upsert(conn, {"guid": "g-remote-new", "data": dict(card)})
with db.get_connection() as c:
    r = c.execute("SELECT status FROM recipients WHERE guid='g-remote-new'").fetchone()
ok("S3a incoming NEW card with 'הסתיים' → 'מושהה'", r and r["status"] == SUSP, str(r and r["status"]))
card2 = dict(card, full_name="כהן אחד", status=ENDED, updated_at="2031-01-01T00:00:00Z")
g_a = db.get_recipient(b)["guid"]
with db.get_connection() as conn:
    sync._apply_rec_upsert(conn, {"guid": g_a, "data": card2})
ok("S3b incoming UPDATE of an existing card with 'הסתיים' → 'מושהה'", raw_status(b) == SUSP, raw_status(b))

# ── S4: מקבל שנמחק ושוחזר — כרטיס ישן עם 'הסתיים' ────────────────────────────────
c_id = db.add_recipient({"full_name": "נמחק ישן", "frequency": "שבועי"})
db.delete_recipient(c_id)
with db.get_connection() as conn:
    row = conn.execute("SELECT guid, card_json FROM deleted_recipients "
                       "WHERE full_name='נמחק ישן'").fetchone()
    import json as _json
    cj = _json.loads(row["card_json"]); cj["status"] = ENDED
    conn.execute("UPDATE deleted_recipients SET card_json=? WHERE guid=?",
                 (_json.dumps(cj, ensure_ascii=False), row["guid"]))
new_id, _msg = db.restore_deleted_recipient(row["guid"])
ok("S4 restored old card with 'הסתיים' comes back 'מושהה'",
   new_id and raw_status(new_id) == SUSP, str(new_id and raw_status(new_id)))

# ── S5: המרה חד-פעמית של DB ישן בהפעלה — אחרי גיבוי, בלי לגעת בחותמת ──────────────
def _raw_insert(name, status, ts="2026-05-05T10:00:00Z", guid=None):
    with db.get_connection() as c:
        c.execute("INSERT INTO recipients (full_name, status, frequency, guid, updated_at) "
                  "VALUES (?,?,?,?,?)", (name, status, "שבועי", guid or ("g-" + name), ts))
        return c.execute("SELECT id FROM recipients WHERE full_name=?", (name,)).fetchone()["id"]

def _safety_files():
    return glob.glob(os.path.join(db.BACKUP_DIR, "safety_*"))

before_files = len(_safety_files())
old1 = _raw_insert("ישן1", ENDED)
old2 = _raw_insert("ישן2", ENDED)
live = _raw_insert("חי1", "פעיל")
sus = _raw_insert("מושהה1", SUSP)
db.init_db()
ok("S5a DB with 'הסתיים' rows: all become 'מושהה'", raw_status(old1) == SUSP and raw_status(old2) == SUSP)
ok("S5b active / suspended rows untouched", raw_status(live) == "פעיל" and raw_status(sus) == SUSP)
with db.get_connection() as c:
    left = c.execute("SELECT COUNT(*) n FROM recipients WHERE status=?", (ENDED,)).fetchone()["n"]
    ts_now = c.execute("SELECT updated_at FROM recipients WHERE id=?", (old1,)).fetchone()["updated_at"]
ok("S5c no 'הסתיים' left in the table", left == 0, str(left))
ok("S5d updated_at untouched (no LWW noise on the other computer)", ts_now == "2026-05-05T10:00:00Z", ts_now)
ok("S5e a safety backup was taken before converting", len(_safety_files()) == before_files + 1,
   f"{before_files}→{len(_safety_files())}")
db.init_db()
ok("S5f second start: idempotent, no extra backup", len(_safety_files()) == before_files + 1)
ok("S5g the column is still called 'status' (no rename)",
   "status" in {r[1] for r in sqlite3.connect(db.DB_PATH).execute("PRAGMA table_info(recipients)")})

# ── S6: מסכים — סינון, כרטיס, ספירה ────────────────────────────────────────────
from main import MainWindow
win = MainWindow()
rt = win.recipients_tab
items = [rt.status_filter.itemText(i) for i in range(rt.status_filter.count())]
ok("S6a list filter: הכל / פעיל / מושהה", items == ["הכל", "פעיל", "מושהה"], str(items))
from tabs.recipients import RecipientDialog
dlg = RecipientDialog(win, db.get_recipient(b))
items = [dlg.f_status.itemText(i) for i in range(dlg.f_status.count())]
ok("S6b card status combo: פעיל / מושהה only", items == ["פעיל", "מושהה"], str(items))
# כרטיס שנטען עם ערך ישן גולמי — חייב להיראות "מושהה", לא "פעיל" (אחרת שמירה תפעיל אותו בטעות)
dlg2 = RecipientDialog(win, dict(db.get_recipient(b), status=ENDED))
ok("S6c a raw legacy 'הסתיים' opens as 'מושהה' (never silently 'פעיל')",
   dlg2.f_status.currentText() == SUSP, dlg2.f_status.currentText())
summ = db.get_summary()
ok("S6d summary: no 'ended' bucket; suspended counts the converted",
   "ended" not in summ and summ["suspended"] >= 3, str({k: summ[k] for k in ("active", "suspended")}))
from utils.ui import STATUS_BADGES
ok("S6e status badges: פעיל / מושהה only", set(STATUS_BADGES) == {"פעיל", SUSP}, str(set(STATUS_BADGES)))

# ── S7: מי מקבל — מקור-אמת יחיד: רק status='פעיל' ───────────────────────────────
weekly = {r["id"] for r in db.get_weekly_list()}
ok("S7a suspended (ex-'הסתיים') is not in the weekly list", old1 not in weekly and b not in weekly)
ok("S7b get_all_recipients('פעיל') excludes them",
   not ({old1, old2, b} & {r["id"] for r in db.get_all_recipients("פעיל")}))
src = open("database.py", encoding="utf-8").read()
vals = set(re.findall(r"status\s*=\s*'([^'?]+)'", src))
# ערכי סטטוס של מקבל בלבד (קמפיינים/משוב משתמשים ב-open/done/sending…)
rec_vals = {v for v in vals if re.search(r"[֐-׿]", v)}
ok("S7c every recipient-status literal in database.py is פעיל/מושהה (or the legacy mapper)",
   rec_vals <= {"פעיל", SUSP, ENDED}, str(rec_vals))
n_ended_sql = len(re.findall(r"status\s*=\s*'" + ENDED + "'", src))
ok("S7d no SQL still counts/filters on status='הסתיים' (except the one-time migration)",
   n_ended_sql <= 1, str(n_ended_sql))
ui_src = open("tabs/recipients.py", encoding="utf-8").read()
ok("S7e tabs/recipients.py has no 'הסתיים'", ENDED not in ui_src)

# ── S8: קמפיין צינתוק — 'done' עדיין "הסתיים" בתצוגה, לא נגענו ─────────────────────
tz_src = open(os.path.join("tabs", "tzintukim.py"), encoding="utf-8").read()
ok("S8 tzintuk campaign status 'done' logic untouched",
   'camp.get("status") == "done"' in tz_src)

print()
if fails:
    print("FAILED:", fails); sys.exit(1)
print("ALL STATUS-SIMPLIFY CHECKS PASS ✓"); sys.exit(0)
